"""Builds rag-workshop.pptx, the slides for the 90-minute workshop.

Run from this folder:  uv run --with python-pptx --with "qrcode[pil]" python build_deck.py
Needs rsvg-convert (brew install librsvg) to turn ../img/*.svg into pictures.
For the PDF, open the .pptx in PowerPoint: File > Export > PDF.
Layout and images follow the KISZ deck from the Brandenburger Digitaltag 2026.
"""
import subprocess
import tempfile
from pathlib import Path

import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
ASSETS = HERE / "assets"
TMP = Path(tempfile.mkdtemp())
REPO = "github.com/aihpi/workshop-rag-basics"

RED, ORANGE, YELLOW, BLUE = "B1063A", "DD6108", "F6A800", "007A9E"
GREY, MUTED, INK, DARK, SAND = "5A6065", "8B9094", "17161C", "30343A", "EDEBE7"
WHITE = "FFFFFF"

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(18288000), Emu(10477500)  # 20 × 11.46 in, as the original
W, H = 20.0, 10477500 / 914400
BLANK = prs.slide_layouts[6]


# --------------------------------------------------------------------------- helpers
def rgb(hex_):
    return RGBColor.from_string(hex_)


def text(slide, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font="Arial", spacing=0, line=1.15, rotation=0):
    """runs: a string, or a list of paragraphs; a paragraph is a string or a list of (text, {overrides})."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.rotation = rotation
    frame = box.text_frame
    frame.word_wrap, frame.vertical_anchor = True, anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for i, para in enumerate([runs] if isinstance(runs, str) else runs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment, p.line_spacing = align, line
        for run_text, opts in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = run_text
            f = r.font
            f.name, f.size = opts.get("font", font), Pt(opts.get("size", size))
            f.bold, f.color.rgb = opts.get("bold", bold), rgb(opts.get("color", color))
            if opts.get("link"):
                r.hyperlink.address = opts["link"]
            if opts.get("spacing", spacing):
                r._r.get_or_add_rPr().set("spc", str(int(opts.get("spacing", spacing) * 100)))
        if isinstance(para, list) and para and para[0][1].get("space_after"):
            p.space_after = Pt(para[0][1]["space_after"])
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb, s.line.width = rgb(line), Pt(1.5)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def poly(slide, points, fill):
    builder = slide.shapes.build_freeform(Inches(points[0][0]), Inches(points[0][1]), scale=Inches(1))
    builder.add_line_segments([(x - points[0][0], y - points[0][1]) for x, y in points[1:]])
    s = builder.convert_to_shape(Inches(points[0][0]), Inches(points[0][1]))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def line(slide, points, color, width=2.0, arrow=False):
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb, c.line.width = rgb(color), Pt(width)
    if arrow:
        c.line._get_or_add_ln().append(c.line._ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))


def image(slide, name, x, y, w=None, h=None):
    path = name if isinstance(name, Path) else ASSETS / name
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                    Inches(w) if w else None, Inches(h) if h else None)


def diagram(name, x, y, w, slide):
    png = TMP / f"{name}.png"
    if not png.exists():
        subprocess.run(["rsvg-convert", "-z", "3", "-b", "white", str(HERE.parent / "img" / f"{name}.svg"), "-o", str(png)], check=True)
    return image(slide, png, x, y, w=w)


def qr(url, name):
    path = TMP / f"{name}.png"
    qrcode.make(url, border=1).save(path)
    return path


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def page_number(slide):
    pass  # numbers are added at the end, once the order is final


def logos(slide):
    image(slide, "logo-bmftr.png", 13.94, 0.35, w=2.62)
    image(slide, "logo-kisz.png", 16.77, 0.43, w=2.8)


def content(title, band, label, logo=True, background=None):
    """White slide with the coloured band on the left, a vertical section label and the title."""
    slide = prs.slides.add_slide(BLANK)
    if background:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(background)
    poly(slide, [(0, 0), (0.94, 0), (0.94, 8.85), (0, H)], band)
    if label:
        width = 0.36 * len(label) + 0.6
        text(slide, 0.47 - width / 2, 6.0 - 0.18, width, 0.36, label, size=15, bold=True, color=WHITE if band != YELLOW else INK,
             spacing=4, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, rotation=270)
    text(slide, 1.16, 0.43, 12.5, 1.4, title.upper(), size=39, bold=True, color=WHITE if background else GREY,
         anchor=MSO_ANCHOR.MIDDLE, line=1.0)
    if logo:
        logos(slide)
    page_number(slide)
    return slide


def divider(number, title, subtitle, fill, ink=WHITE):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(fill)
    text(slide, 1.16, 2.85, 12, 0.5, number, size=22, bold=True, color=ink, spacing=3)
    text(slide, 1.16, 3.45, 17, 2.3, title.upper(), size=60, bold=True, color=ink, line=0.95, anchor=MSO_ANCHOR.BOTTOM)
    rect(slide, 1.16, 6.0, 1.75, 0.08, ink)
    text(slide, 1.16, 6.45, 16, 0.6, subtitle, size=24, color=ink)
    poly(slide, [(0, 8.85), (4.65, H), (0, H)], WHITE)
    image(slide, "logo-bmftr.png", 0.05, 9.75, w=2.62)
    page_number(slide)
    return slide


def numbered(slide, x, y, w, items, colour=YELLOW, size=18, gap=0.72):
    """Numbered squares; an item starting with "Bonus" gets a + and a paler square."""
    number = 0
    for i, item in enumerate(items, start=1):
        bonus = item.startswith("Bonus")
        number += not bonus
        rect(slide, x, y + (i - 1) * gap, 0.42, 0.42, SAND if bonus else colour)
        text(slide, x, y + (i - 1) * gap, 0.42, 0.42, "+" if bonus else str(number), size=16, bold=True,
             color=WHITE if colour != YELLOW and not bonus else INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, x + 0.65, y + (i - 1) * gap + 0.03, w - 0.65, gap, item, size=size)


def kicker(slide, x, y, w, words, color=MUTED):
    text(slide, x, y, w, 0.4, words.upper(), size=15, bold=True, color=color, spacing=3)


def chat(slide, x, y, w, h, exchange):
    """A chat window drawn from shapes: [(who, text, extras)], who is 'you', 'bot' or 'step'."""
    rect(slide, x, y, w, h, "F4F3F1", shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.04
    text(slide, x + 0.35, y + 0.25, w - 0.7, 0.35, "localhost:8000", size=13, color=MUTED, font="Courier New")
    cy = y + 0.85
    for who, body, height in exchange:
        if who == "you":
            bw = min(w - 1.6, 0.13 * len(body) + 0.6)
            rect(slide, x + w - 0.35 - bw, cy, bw, height, "E2E0DC", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(slide, x + w - 0.2 - bw, cy + 0.12, bw - 0.3, height - 0.2, body, size=15)
        elif who == "step":
            text(slide, x + 0.35, cy, w - 0.7, height, [[("✓ ", {"color": ORANGE, "bold": True}), (body, {"color": GREY})]], size=15)
        else:
            rect(slide, x + 0.35, cy, w - 0.7, height, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.08
            text(slide, x + 0.6, cy + 0.15, w - 1.2, height - 0.25, body, size=15)
        cy += height + 0.25


def notebook(number, title, minutes, learn, turn, exchange):
    slide = content(title, YELLOW, "HANDS-ON")
    kicker(slide, 1.16, 1.85, 9, f"Notebook {number} · {minutes} min", ORANGE)
    kicker(slide, 1.16, 2.65, 9, "In this notebook")
    text(slide, 1.16, 3.15, 9.4, 2.4, [[(f"•  {item}", {"space_after": 8})] for item in learn], size=19)
    kicker(slide, 1.16, 5.55, 9, "Your turn · in the chatbot")
    numbered(slide, 1.16, 6.05, 9.4, turn, gap=0.75)
    text(slide, 1.16, 6.05 + 0.75 * len(turn) + 0.15, 9.4, 0.4, "Solutions are folded away under each exercise.",
         size=15, color=MUTED)
    chat(slide, 11.2, 2.3, 8.3, 1.1 + sum(h + 0.25 for _, _, h in exchange), exchange)
    return slide


# ============================================================================ main path, slides 1 to 20
# Sessions: S0 welcome and setup · S1 chat · S2 documents · S3 search · S4 RAG and faithfulness · S5 your PDF.

def vote(title, intro, questions, reveal=None):
    """A checkpoint on grey. questions: [(kicker, question, options)]; reveal: [(letter, why)] or None."""
    s = content(title, YELLOW, "CHECKPOINT", logo=False, background=GREY)
    if reveal is None:
        rect(s, 13.0, 0.85, 4.9, 1.1, YELLOW)
        text(s, 13.25, 0.85, 4.5, 1.1, intro, size=17, anchor=MSO_ANCHOR.MIDDLE)
    width = 17.7 / len(questions)
    for i, (where, question, options) in enumerate(questions):
        x = 1.16 + i * width
        rect(s, x, 2.5, width - 0.6, 0.08, YELLOW)
        text(s, x, 2.85, width - 0.6, 0.4, where.upper(), size=15, bold=True, color="D0D3D6", spacing=3)
        text(s, x, 3.35, width - 0.7, 1.4, question, size=26, bold=True, color=WHITE, line=1.05)
        if reveal is None:
            text(s, x, 5.0, width - 0.7, 4, [[(f"{'ABCD'[j]}   {o}", {"space_after": 12})] for j, o in enumerate(options)],
                 size=21, color=WHITE, line=1.2)
            continue
        letter, why = reveal[i]
        rect(s, x, 4.95, width - 0.6, 3.6, WHITE)
        text(s, x + 0.3, 5.15, width - 1.2, 0.9, [[(letter, {"color": RED, "size": 26, "bold": True}),
                                                    (f"   {options['ABCD'.index(letter)]}", {"bold": True})]], size=19)
        text(s, x + 0.3, 6.3, width - 1.2, 2.2, why, size=19, line=1.25)
    return s


# ---------------------------------------------------------------------------- 1 title
s = prs.slides.add_slide(BLANK)
text(s, 1.67, 2.24, 15, 0.4, "[EVENT]", size=19.5, bold=True, color="DC640D", spacing=4)
text(s, 1.67, 2.82, 15, 1.1, "HANDS-ON RAG", size=69, bold=True, color=GREY)
text(s, 1.67, 4.11, 15, 0.7, "Build your own document Q&A system", size=34.5, color=RED)
rect(s, 1.67, 5.2, 1.25, 0.08, RED)
text(s, 1.67, 5.76, 14, 1.1, [[("Felix Boelter  |  [Co-moderator]", {"size": 22.5})],
                              [("AI Service Centre Berlin-Brandenburg  |  [Date]", {"size": 22.5, "color": GREY})]])
image(s, "logo-kisz.png", 16.45, 0.74, w=2.8)
image(s, "logo-bmftr.png", 16.41, 1.98, w=2.8)
image(s, "hpi-tagline.png", 0.76, 10.23, w=2.92)
image(s, "hpi-building.png", 9.72, 6.72, w=11.72)
notes(s, "S0, 15 minutes with the install. 90 minutes, mostly in the notebooks. The participant list is going round.")

# ---------------------------------------------------------------------------- 2 participant list, start the install
s = content("Participant list", MUTED, None)
text(s, 1.16, 2.6, 10, 1.2, "The list is going round now. Please sign in while everyone is still seated.", size=24)
for i, head in enumerate(["Name", "Institution", "Email"]):
    rect(s, 1.16 + i * 3.13, 4.3, 3.1, 0.75, ORANGE)
    text(s, 1.36 + i * 3.13, 4.3, 2.8, 0.75, head, size=19, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    for row in range(3):
        rect(s, 1.16 + i * 3.13, 5.08 + row * 0.6, 3.1, 0.57, SAND)
text(s, 1.16, 7.3, 9.4, 1.0, "The printout stays with the workshop lead. Please hand it back at the end.", size=17, color=GREY)
rect(s, 11.4, 2.6, 8.1, 6.2, YELLOW)
text(s, 11.8, 2.9, 7.4, 0.5, "START THE INSTALL NOW", size=17, bold=True, spacing=3)
text(s, 11.8, 3.5, 7.4, 1.2, "It downloads about 2 GB, so let it run while we talk.", size=20, line=1.2)
text(s, 11.8, 4.75, 7.4, 2.6, "git clone https://github.com/\n    aihpi/workshop-rag-basics.git\ncd workshop-rag-basics/workshop\nuv sync",
     size=17, font="Courier New", line=1.3)
image(s, qr("https://" + REPO, "qr-repo"), 11.8, 7.0, w=1.5)
text(s, 13.5, 7.35, 5.8, 0.9, "Needs uv and Git, from the\nworkshop-getting-started guide.", size=15, line=1.2)
notes(s, "Pass the list round now, not at the end. Get everyone to start git clone and uv sync now: it's about 2 GB "
         "on the workshop Wi-Fi, and it runs while we talk through the next three slides. Collect the list at the end.")

# ---------------------------------------------------------------------------- 3 destination and agenda
s = content("By 1:30 you have", MUTED, None)
for i, line_ in enumerate(["A chatbot on your laptop that answers from your own PDF.",
                           "Every claim links to the page it came from, with a score that says how well it is backed.",
                           "A list of what to check before you trust it."]):
    rect(s, 1.16, 2.3 + i * 1.25, 0.55, 0.55, YELLOW)
    text(s, 1.16, 2.3 + i * 1.25, 0.55, 0.55, str(i + 1), size=20, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.0, 2.3 + i * 1.25, 16.5, 0.9, line_, size=26, anchor=MSO_ANCHOR.MIDDLE)
kicker(s, 1.16, 6.5, 9, "Agenda")
AGENDA = [(0, 15, RED, "Welcome · setup"), (15, 13, YELLOW, "1 · Chat"), (28, 13, YELLOW, "2 · Documents"),
          (41, 13, YELLOW, "3 · Search"), (54, 16, YELLOW, "4 · RAG + score"), (70, 18, YELLOW, "5 · Your PDF · transfer"),
          (88, 2, MUTED, "")]
per_minute = 17.7 / 90
for start, minutes, colour, label in AGENDA:
    x, w = 1.16 + start * per_minute, minutes * per_minute - 0.05
    rect(s, x, 7.55, w, 1.0, colour)
    if label:
        text(s, x + 0.12, 7.55, w - 0.2, 1.0, label, size=15, bold=True, color=WHITE if colour == RED else INK,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x, 7.05, w, 0.4, f"{start // 60}:{start % 60:02d}", size=15, bold=True, color=GREY)
text(s, 1.16, 8.85, 17.7, 0.5, "Every notebook block: a minute of slides, then the notebook, then a quick check.",
     size=17, color=GREY)
notes(s, "Read the three lines aloud. Ask: is this what you came for? Write anything that doesn't fit on the whiteboard; "
         "we come back to it at the end (slide 18).")

# ---------------------------------------------------------------------------- 4 who we are, one slide
s = content("AI Service Centre\nBerlin-Brandenburg", RED, None)
for i, (colour, head, body) in enumerate([
        (YELLOW, "Free workshops, talks and MOOCs", "Like today, and online at openHPI"),
        (DARK, "Free GPU infrastructure", "For research and teaching, request access at aisc.hpi.de"),
        (ORANGE, "Office hours and pilot projects", "For your own project: more at the end")]):
    y = 2.4 + i * 2.2
    rect(s, 1.16, y, 0.18, 1.45, colour)
    text(s, 1.65, y - 0.02, 10.5, 0.6, head, size=28, bold=True)
    text(s, 1.65, y + 0.68, 10.5, 0.6, body, size=21, color=GREY)
f = 4.4 / 5.21  # the map from the appendix slide, smaller
image(s, "map-germany.png", 13.6, 2.3, w=4.4)
image(s, "map-states.png", 13.6, 2.3 + 0.13 * f, w=5.19 * f)
for colour, dots in [("1BABE3", [(7.75, 4.08), (7.8, 4.88), (7.5, 5.2)]), (ORANGE, [(9.55, 4.29)]),
                     ("E9358D", [(6.37, 4.99), (6.42, 5.25), (5.92, 5.59), (5.74, 5.7), (6.52, 5.97)]), ("8A6AF4", [(6.97, 6.24)])]:
    for dx, dy in dots:
        x, y = 13.6 + (dx - 5.46) * f, 2.3 + (dy - 1.85) * f
        size = 0.22 if colour == ORANGE else 0.11
        rect(s, x - size / 2, y - size / 2, size, size, colour, shape=MSO_SHAPE.OVAL)
text(s, 13.6, 8.35, 4.4, 0.5, "Four AI service centres, funded by the BMFTR", size=15, color=GREY, align=PP_ALIGN.CENTER)
notes(s, "30 seconds. The four centres, our offers in one line each. Details are in the appendix; office hours and "
         "pilot projects come back on slide 19.")

# ---------------------------------------------------------------------------- 5 architecture
s = content("Architecture", BLUE, "BASICS")
rect(s, 1.32, 2.2, 4.05, 6.1, "EAF4F8", line="9CCBDD")
rect(s, 5.52, 2.2, 9.63, 6.1, "FDF1E9", line="F2B994")
for x, n, colour, head, sub in [(1.55, "1", BLUE, "INGESTION", "once · notebooks 2 + 3"),
                                (5.75, "2", ORANGE, "INFERENCE", "every question · notebooks 3 + 4")]:
    rect(s, x, 2.42, 0.42, 0.42, colour)
    text(s, x, 2.42, 0.42, 0.42, n, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.6, 2.38, 6, 0.45, head, size=19, bold=True, color=GREY, spacing=3)
    text(s, x + 0.6, 2.82, 7, 0.4, sub, size=15, color=MUTED)
boxes = {  # x, y, w, fill, ink, title, subtitle
    "docs": (2.0, 3.36, 3.15, "E7E5E1", INK, "Documents", "PDFs"),
    "embed": (2.0, 5.05, 5.4, "E7E5E1", INK, "Embedding model", "Text in, vector out"),
    "db": (2.0, 6.72, 5.4, BLUE, WHITE, "Vector database", "Index for similarity search"),
    "ask": (8.23, 3.36, 3.17, ORANGE, WHITE, "Question", "asked by the user"),
    "hits": (8.23, 6.72, 3.17, "E7E5E1", INK, "Relevant passages", "go into the prompt"),
    "llm": (11.98, 5.05, 3.13, "E7E5E1", INK, "LLM", "language model"),
    "answer": (15.69, 5.05, 3.17, ORANGE, WHITE, "Answer", "with sources"),
}
for x, y, w, fill, ink, head, sub in boxes.values():
    rect(s, x, y, w, 1.29, fill)
    text(s, x + 0.27, y + 0.22, w - 0.4, 0.45, head, size=20, bold=True, color=ink)
    text(s, x + 0.27, y + 0.7, w - 0.4, 0.4, sub, size=16, color=ink)
line(s, [(3.57, 4.65), (3.57, 5.0)], BLUE, 3, arrow=True)
line(s, [(3.57, 6.34), (3.57, 6.67)], BLUE, 3, arrow=True)
line(s, [(8.35, 4.77), (7.5, 5.62)], ORANGE, 3, arrow=True)
line(s, [(6.04, 6.34), (6.04, 6.67)], ORANGE, 3, arrow=True)
line(s, [(7.4, 7.36), (8.18, 7.36)], ORANGE, 3, arrow=True)
line(s, [(11.25, 4.77), (11.92, 5.47)], ORANGE, 3, arrow=True)
line(s, [(11.25, 6.67), (11.92, 5.97)], ORANGE, 3, arrow=True)
line(s, [(15.14, 5.7), (15.64, 5.7)], ORANGE, 3, arrow=True)
text(s, 1.16, 8.85, 8.4, 2.0, [[("Once, ahead of time, ", {"bold": True, "color": BLUE}),
                                ("the documents are read and cut into chunks. The embedding model turns every chunk into a "
                                 "vector, a list of numbers that captures its meaning. The vectors go into the vector database.", {})]], size=17, line=1.25)
text(s, 10.3, 8.85, 8.6, 2.0, [[("For every question, ", {"bold": True, "color": ORANGE}),
                                ("the same model turns the question into a vector and searches the same database. The closest "
                                 "chunks go to the language model with the question, and it answers with sources.", {})]], size=17, line=1.25)
notes(s, "Every box on this picture is one notebook. Ingestion happens once, inference with every question; the same "
         "embedding model in both. Ask: do we agree on the route?")

# ---------------------------------------------------------------------------- 6 setup
s = content("Setup", YELLOW, "HANDS-ON")
steps = [("Install (started on slide 2)", "git clone https://github.com/aihpi/workshop-rag-basics.git\ncd workshop-rag-basics/workshop\nuv sync"),
         ("Add the key", "cp .env.example .env    # then paste the key from the whiteboard"),
         ("Check, and fetch Docling's models", "uv run python check_setup.py"), ("Start", "uv run jupyter lab")]
y = 2.15
for i, (head, code) in enumerate(steps):
    rect(s, 1.16, y, 0.42, 0.42, YELLOW)
    text(s, 1.16, y, 0.42, 0.42, str(i + 1), size=16, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.85, y - 0.02, 9, 0.45, head, size=19, bold=True)
    text(s, 1.85, y + 0.45, 9.8, 1.2, code, size=15, font="Courier New", color=GREY, line=1.2)
    y += 1.3 + 0.27 * code.count("\n")
text(s, 1.16, y + 0.2, 10, 0.8, [[("Done when check_setup.py prints ", {}), ("You're ready for the workshop.", {"bold": True})]],
     size=18)
diagram("live-reload", 12.0, 2.2, 7.6, s)
text(s, 12.0, 4.3, 7.4, 1.8, "Cells that start with %%writefile app.py write the chatbot. Chainlit notices the change and "
     "reloads it in your browser at localhost:8000.", size=18, line=1.25)
for i, (head, value) in enumerate([("Chat model", "gemma-4-31b"), ("Embedding model", "octen-embedding-8b"),
                                   ("Runs on", "our own infrastructure in Potsdam")]):
    kicker(s, 12.0, 6.55 + i * 1.15, 7, head)
    text(s, 12.0, 6.95 + i * 1.15, 7.4, 0.45, value, size=19, font="Courier New" if i < 2 else "Arial")
notes(s, "Key on the whiteboard now. check_setup.py downloads about 500 MB of Docling models the first time; notebook 1 "
         "works before that finishes. Anyone whose uv sync is still running: raise a hand. Ask: everyone sees "
         "'You're ready'?")

# ---------------------------------------------------------------------------- 7 notebook 1
s = notebook(1, "Talking to an LLM", 6,
             ["Connect to the model through the gateway", "Ask about a detail from one of our papers",
              "Build a chatbot that remembers and streams"],
             ["Personality: change the system prompt", "Bonus: stopwatch, you'll want it in notebook 4"],
             [("you", "In Kage et al. (2018), what were the lifetime-encoded beads loaded with?", 1.0),
              ("step", "The model answers. Is it right?", 0.45)])
notes(s, "S1. Don't say whether the answer is right. They decide in five minutes.")

# ---------------------------------------------------------------------------- 8 the vote
s = vote("Was it right?", "Hands up.\nRemember your letter.",
         [("Notebook 1", "The model answered the Kage question. Was it ...", ["right", "wrong", "partly right", "can't tell"])])
notes(s, "Count hands per letter, write the numbers on the whiteboard. No reveal: notebook 4 tells us.")

# ---------------------------------------------------------------------------- 9 why
s = content("Half right, and you can't\ntell which half", RED, "MOTIVATION")
cols = [(ORANGE, "No access to\nyour documents",
         "Language models are trained mostly on public data. Your lab reports and paywalled papers are unknown to them."),
        (RED, "Similar instead\nof right",
         "When the right information is missing, the model falls back on related material: the rules for no-parking "
         "zones answer a question about no-stopping zones."),
        (YELLOW, "RAG decides what\nthe answer is built on",
         "Retrieval-augmented generation sets which documents an answer may use, and cites the page.")]
for i, (colour, head, body) in enumerate(cols):
    x = 1.16 + i * 6.12
    rect(s, x, 2.5, 5.46, 0.08, colour)
    text(s, x, 2.85, 5.5, 0.4, f"0{i + 1}", size=19, bold=True, color=MUTED, spacing=3)
    text(s, x, 3.35, 5.5, 1.2, head, size=28, bold=True, line=1.05)
    text(s, x, 4.75, 5.3, 4.5, body, size=22, line=1.3)
notes(s, "Most of you voted C or D. That is the honest answer, and it is the problem. Notebook 4 tells us which half.")

# ---------------------------------------------------------------------------- 10 notebook 2
s = content("Reading documents", YELLOW, "HANDS-ON")
kicker(s, 1.16, 1.85, 9, "Notebook 2 · 7 min", ORANGE)
diagram("docling", 1.16, 2.45, 8.6, s)
text(s, 1.16, 6.2, 8.5, 1.4, "A PDF stores where to draw each letter. Docling finds headings, paragraphs and tables, "
     "and writes Markdown.", size=19, line=1.25)
diagram("chunking", 10.6, 2.45, 8.6, s)
text(s, 10.6, 6.2, 8.6, 1.4, "We cut chunks of 1000 characters; a new one starts every 800, so neighbours share 200.",
     size=19, line=1.25)
kicker(s, 1.16, 8.05, 9, "Your turn · in the chatbot")
numbered(s, 1.16, 8.55, 17, ["Where are the cuts? Try 300 and 3000", "Bonus: only the tables"], gap=0.7)
notes(s, "S2. Docling needs about a minute for the three papers; talk through the Markdown while it runs. "
         "GrundschutzKI cuts at headings instead; that is in the appendix.")

# ---------------------------------------------------------------------------- 11–12 notebook 2 checkpoint
NB2 = [("Notebook 2", "What happened to the tables when Docling read the papers?",
        ["They were dropped", "They became loose lines of numbers", "They stayed tables, in Markdown"]),
       ("Notebook 2 · your turn", "What happens with smaller chunks, 300 characters instead of 1000?",
        ["Fewer chunks, and tables stay whole", "More chunks, and more of them cut through a table", "Nothing changes"])]
s = vote("What did you find?", "Vote by show of hands.\nAnswers on the next slide.", NB2)
notes(s, "Read out, hands up.")
s = vote("What we found", None, NB2, reveal=[
    ("C", "All 5 tables in the three papers came through with their rows and columns, searchable like any other text."),
    ("B", "628 instead of 197 chunks, and 17 instead of 5 end in the middle of a table row. At 3000: 69 chunks, none.")])
rect(s, 1.16, 9.0, 17.7, 1.1, WHITE)
text(s, 1.5, 9.0, 3, 1.1, "NEXT", size=15, bold=True, color=MUTED, spacing=3, anchor=MSO_ANCHOR.MIDDLE)
text(s, 3.4, 9.0, 15.2, 1.1, "We have good chunks. How do we find the right ones for a question?", size=20, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Reveal. Anyone see something different? The chunk numbers are measured on the three papers.")

# ---------------------------------------------------------------------------- 13 notebook 3
s = notebook(3, "Search", 8,
             ["Embed sentences and compare how similar they are", "Store all chunks in Qdrant, in memory",
              "Turn the bot into a search engine, no LLM yet"],
             ["Fewer, better results: pick a score cutoff", "Bonus: ask in German"],
             [("you", "What were the lifetime-encoded beads loaded with?", 0.62),
              ("bot", "Kage_2018_SciReports.pdf, p. 2  (score 0.73)\n› … loaded with different organic fluorophores …\n\n"
                      "Kage_2018_SciReports.pdf, p. 3  (score 0.64)\n› Table 1. Lifetime codes and respective luminophores …", 2.55),
              ("you", "What is the capital of France?", 0.62),
              ("bot", "Kage_2018_SciReports.pdf, p. 6  (score 0.17) …", 0.75)])
notes(s, "S3. Make them ask the France question before you explain anything.")

# ---------------------------------------------------------------------------- 14 why France
s = content("Why did France get a hit?", BLUE, "BASICS")
diagram("embeddings", 1.16, 2.1, 7.4, s)
diagram("qdrant", 8.95, 2.1, 10.6, s)
for i, line_ in enumerate(["An embedding model turns text into a list of numbers, 4096 with ours. Similar meaning, similar direction.",
                           "Search returns the closest chunks, however far away they are.",
                           "Qdrant stores one point per chunk: the vector, and a payload with text, file and page."]):
    text(s, 1.16, 6.85 + i * 1.0, 17.7, 0.9, [[(f"{i + 1}   ", {"bold": True, "color": BLUE}), (line_, {})]], size=20)
notes(s, "The cat and kitten numbers are the ones they just saw. Ask: do we know what we've built so far?")

# ---------------------------------------------------------------------------- 15 notebook 4
s = content("Retrieval-augmented generation", YELLOW, "HANDS-ON")
kicker(s, 1.16, 1.85, 9, "Notebook 4 · 11 min", ORANGE)
for i, (word, rest) in enumerate([("Retrieve", "the 5 closest chunks"), ("Augment", "the prompt with them, numbered [1] to [5]"),
                                  ("Generate", "the answer, citing those numbers")]):
    text(s, 1.16, 2.55 + i * 1.05, 8.6, 1.0, [[(word.upper(), {"bold": True, "color": BLUE, "spacing": 3, "size": 17})],
                                              [(rest, {"size": 20})]])
text(s, 1.16, 5.75, 8.6, 1.0, "Every [n] leads to a chunk, and every chunk remembers its file and page.", size=20, line=1.25)
kicker(s, 1.16, 7.15, 9, "Your turn")
numbered(s, 1.16, 7.65, 8.8, ["How much context? Score limit 2 and limit 15", "Bonus: strict or not",
                              "Bonus: show the evidence; follow-up questions"], gap=0.72)
diagram("rag-prompt", 10.5, 2.0, 9.0, s)
chat(s, 10.5, 5.7, 9.0, 1.1 + 0.62 + 0.25 + 1.9 + 0.25,
     [("you", "What were the beads in Kage et al. loaded with?", 0.62),
      ("bot", "PMMA beads stained with organic dyes from PolyAn GmbH [1], and melamine beads loaded with "
              "CdSe/CdS/ZnS quantum dots [1], [2].\n\nSources:  Kage_2018_SciReports, p. 2", 1.9)])
notes(s, "S4. Same question as notebook 1. Let them click a source. Strict-or-not is a bonus today; it comes back in the "
         "transfer round.")

# ---------------------------------------------------------------------------- 16 which half
s = content("Which half?", YELLOW, "CHECKPOINT", logo=False, background=GREY)
rect(s, 1.16, 2.3, 10.4, 5.9, WHITE)
text(s, 1.5, 2.5, 9.8, 0.4, "THE NOTEBOOK 1 ANSWER, JUDGED IN NOTEBOOK 4 (ONE RUN)", size=14, bold=True, color=MUTED, spacing=2)
JUDGED = [(True, "The beads were loaded with quantum dots (QDs)."),
          (False, "Three different types of quantum dots, to create distinct lifetime signatures."),
          (False, "One type had a short lifetime of about 10 ns."),
          (False, "One type had a medium lifetime of about 25 ns."),
          (False, "One type had a long lifetime of about 50 ns."),
          (False, "The beads were used as fiducial markers to correct for drift during imaging.")]
text(s, 1.5, 3.05, 9.8, 5.6, [[("✓  " if ok else "✗  ", {"bold": True, "color": "2F9E62" if ok else RED}), (claim, {})]
                              for ok, claim in JUDGED], size=18, line=1.45)
text(s, 1.5, 7.2, 9.8, 0.7, [[("faithfulness  ", {"color": GREY}), ("0.17", {"bold": True, "size": 26})]], size=20,
     font="Courier New")
text(s, 12.3, 2.3, 7.2, 1.6, "Which part of the system knows the page number of a chunk?", size=26, bold=True,
     color=WHITE, line=1.1)
text(s, 12.3, 4.4, 7.2, 3, [[(f"{l}   {o}", {"space_after": 12})] for l, o in
                            [("A", "the LLM"), ("B", "the payload in Qdrant"), ("C", "the embedding model")]], size=22, color=WHITE)
notes(s, "Go back to the vote on the whiteboard. The ✗ lines are the half nobody could see; one claim (quantum dots) "
         "was right. Their own run will show different claims and a different score. Then the question: B, written "
         "at ingestion from Docling's page numbers (notebook 3, slide 14). Ask: do we see how far we've come?")

# ---------------------------------------------------------------------------- 17 notebook 5
s = content("Putting it together", YELLOW, "HANDS-ON")
kicker(s, 1.16, 1.85, 9, "Notebook 5 · 6 min", ORANGE)
diagram("overview", 1.16, 2.45, 10.3, s)
kicker(s, 12.3, 2.45, 7, "Your turn · in the chatbot")
numbered(s, 12.3, 2.95, 7.2, ["Your own document: attach a PDF, ask, check the page and the score",
                              "Bonus: chat with this document only"], gap=1.25)
pieces = [("Read a PDF with Docling", "2"), ("Cut into chunks", "2"), ("Turn text into vectors", "3"),
          ("Store and search in Qdrant", "3"), ("Answer from the chunks", "4"), ("Cite the page", "4"),
          ("Score the answer", "4"), ("Your own PDF", "5")]
kicker(s, 1.16, 6.6, 9, "Piece · notebook")
for i, (piece, nb) in enumerate(pieces):
    x, y = 1.16 + (i % 4) * 4.45, 7.1 + (i // 4) * 0.95
    rect(s, x, y, 4.25, 0.75, SAND)
    text(s, x + 0.2, y, 3.4, 0.75, piece, size=16, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 3.55, y, 0.5, 0.75, nb, size=20, bold=True, color=ORANGE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.16, 9.25, 17.7, 0.5, [[("Your PDF goes to the AISC gateway in Potsdam. ", {"bold": True}),
                                 ("Test or public documents only.", {})]], size=18, color=RED)
notes(s, "S5. The score is already in the app. Expect at least one answer that is worse than on the papers; that's the "
         "point. A scanned PDF has no text layer, and with OCR off the bot gets nothing to answer from.")

# ---------------------------------------------------------------------------- 18 transfer
s = content("Your collection", YELLOW, "TRANSFER")
for i, question in enumerate(["Which collection would you point this at?", "What do you expect to break first, and why?"]):
    rect(s, 1.16 + i * 9.0, 2.4, 8.6, 3.6, SAND)
    text(s, 1.6 + i * 9.0, 2.75, 7.8, 0.5, f"QUESTION {i + 1}", size=15, bold=True, color=MUTED, spacing=3)
    text(s, 1.6 + i * 9.0, 3.35, 7.8, 2.4, question, size=32, bold=True, line=1.1)
text(s, 1.16, 6.65, 17.7, 0.6, "One sentence each.", size=22, color=GREY)
rect(s, 1.16, 8.1, 17.7, 1.3, WHITE, line="D0D3D6")
text(s, 1.5, 8.1, 3, 1.3, "OFTEN", size=15, bold=True, color=MUTED, spacing=3, anchor=MSO_ANCHOR.MIDDLE)
text(s, 3.4, 8.1, 15.2, 1.3, "Scanned PDFs  ·  tables  ·  questions across several documents", size=24, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "One sentence each; above 12 people, tables of four, one sentence per table. Show the OFTEN line only after "
         "the round. Close the loop with the whiteboard list from slide 3. Ask: can you repeat it on your own? "
         "If someone asks why not put all PDFs in the prompt: cost per question, page citations, and collections "
         "larger than any context window.")

# ---------------------------------------------------------------------------- 19 where to take this
s = content("Where to take this", BLUE, "NEXT STEPS")
rows = [("Today", "The template"),
        ("Fixed 1000-character chunks", "Structure-aware chunkers, chosen per data source"),
        ("Qdrant in memory, rebuilt on every start", "Qdrant server, ingested once, updated when files change"),
        ("Vector search only", "Hybrid search: vectors plus keyword matching"),
        ("Faithfulness, one answer at a time", "An evaluation app that compares setups")]
for i, (today, template) in enumerate(rows):
    y = 2.15 + i * 0.86
    fill = BLUE if i == 0 else ("F2F1EE" if i % 2 else WHITE)
    rect(s, 1.16, y, 17.7, 0.82, fill)
    ink, b = (WHITE, True) if i == 0 else (INK, False)
    text(s, 1.45, y, 6.8, 0.82, today, size=18, bold=b, color=ink, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.6, y, 10.1, 0.82, template, size=18, bold=b, color=ink, anchor=MSO_ANCHOR.MIDDLE)
for i, (code, label, link) in enumerate([(qr("https://github.com/aihpi/pilotproject-rag-template", "qr-template"), "The template",
                                          "aihpi/pilotproject-rag-template on GitHub"),
                                         ("qr-office-hours.png", "Office hours", "Stuck on your own collection? Book a slot."),
                                         ("qr-apply.png", "Pilot projects", "Build it with us; applications every three months.")]):
    x = 1.16 + i * 6.0
    image(s, code, x, 6.85, w=1.9)
    text(s, x + 2.15, 7.0, 3.7, 0.5, label, size=19, bold=True)
    text(s, x + 2.15, 7.55, 3.7, 1.3, link, size=15, color=GREY, line=1.2)
notes(s, "The template is where a real project starts. Office hours if you get stuck on your collection. SENTRA and "
         "GrundschutzKI (appendix) are two pilot projects built this way.")

# ---------------------------------------------------------------------------- 20 feedback
s = prs.slides.add_slide(BLANK)
text(s, 1.67, 1.85, 13.5, 1.2, "YOUR FEEDBACK COUNTS", size=60, bold=True, color=GREY)
text(s, 1.67, 3.1, 9.5, 1.2, "We shape our next workshops around your answers. Two minutes is enough.", size=24, line=1.25)
rect(s, 1.67, 4.55, 1.25, 0.08, RED)
image(s, "qr-feedback.png", 1.67, 4.97, w=2.92)
text(s, 5.05, 5.6, 6, 2.2, [[("kisz@hpi.de", {"link": "mailto:kisz@hpi.de"})], [("hpi.de/kisz", {"link": "https://hpi.de/kisz"})],
                            [("Felix Boelter  |  [Co-moderator]", {"color": GREY, "size": 17})]], size=22, line=1.4)
rect(s, 1.67, 8.3, 7.4, 0.02, "D0D3D6")
text(s, 1.67, 8.5, 8, 0.4, "KEEP GOING", size=15, bold=True, color=RED, spacing=3)
text(s, 1.67, 8.95, 8.5, 1.0, [[(REPO, {"link": "https://" + REPO}), ("  ·  this workshop", {"color": GREY})],
                               [("github.com/aihpi/pilotproject-rag-template", {"link": "https://github.com/aihpi/pilotproject-rag-template"}),
                                ("  ·  template", {"color": GREY})]], size=15, line=1.3)
image(s, "logo-kisz.png", 16.45, 0.74, w=2.8)
image(s, "logo-bmftr.png", 16.41, 1.98, w=2.8)
image(s, "hpi-tagline.png", 0.76, 10.23, w=2.92)
image(s, "hpi-building.png", 9.72, 6.72, w=11.72)
notes(s, "Leave the QR code up while questions come in. Contact: kisz@hpi.de.")


# ============================================================================ appendix
s = divider("", "Appendix", "Background for questions. Not part of the 90 minutes.", GREY)
notes(s, "Everything after this slide is for questions and for reading afterwards.")

s = content("Four AI service centres", BLUE, "AI SERVICE CENTRES")
image(s, "qr-kisz.png", 1.25, 2.75, w=2.85)
text(s, 1.25, 5.65, 2.85, 0.35, "hpi.de/ki-servicezentrum", size=13, color=GREY, align=PP_ALIGN.CENTER)
image(s, "map-germany.png", 5.46, 1.85, w=5.21)
image(s, "map-states.png", 5.46, 1.98, w=5.19)
centres = [  # colour, dots on the map, path to the label, logo, place, focus
    ("1BABE3", [(7.75, 4.08), (7.8, 4.88), (7.5, 5.2)], [(7.99, 4.88), (8.47, 4.88), (10.82, 2.8), (13.39, 2.8)],
     "logo-kisski.png", "Göttingen | Kassel | Hannover", "Sensitive and critical infrastructure, medicine & energy"),
    (ORANGE, [(9.55, 4.29)], [(9.55, 4.29), (13.39, 4.29)],
     "logo-kisz-small.png", "Potsdam | Berlin", "Education and consulting, AI in business & society"),
    ("E9358D", [(6.37, 4.99), (6.42, 5.25), (5.92, 5.59), (5.74, 5.7), (6.52, 5.97)], [(6.45, 5.05), (7.17, 5.52), (13.39, 5.52)],
     "logo-westai.png", "Bonn | St. Augustin | Aachen | Jülich | Dortmund", "AI hardware, consulting and training; multimodal and transferable models"),
    ("8A6AF4", [(6.97, 6.24)], [(7.17, 6.24), (9.6, 6.24), (10.8, 6.8), (13.39, 6.8)],
     "logo-hessian.png", "Darmstadt", "Explainability, generalisability and contextual adaptation"),
]
for i, (colour, dots, path, logo, place, focus) in enumerate(centres):
    for dx, dy in dots:
        rect(s, dx - 0.06, dy - 0.06, 0.12, 0.12, colour, shape=MSO_SHAPE.OVAL)
    line(s, path, colour, 2.25)
    rect(s, path[-1][0] - 0.1, path[-1][1] - 0.1, 0.2, 0.2, colour, shape=MSO_SHAPE.OVAL)
    top = [2.35, 3.8, 5.05, 6.4][i]
    image(s, logo, 11.85, top, w=1.35)
    text(s, 13.7, top - 0.05, 6.2, 0.4, place, size=19, bold=True, color=colour)
    text(s, 13.7, top + 0.4, 6.0, 0.8, focus, size=16)
image(s, "logo-centres.png", 6.01, 9.47, w=4.48)
text(s, 15.2, 8.9, 4.6, 1.8, [[("Goal: ", {"bold": True}), ("lower the barriers to putting AI to work in society and business", {"bold": True})]], size=18)
notes(s, "Four centres across Germany. Our site in Potsdam and Berlin covers education, consulting "
         "and the use of AI in business and society.")

s = content("Education", YELLOW, None, logo=False)
image(s, "qr-newsletter.png", 1.16, 2.4, w=2.85)
text(s, 1.16, 5.3, 2.85, 0.35, "Newsletter", size=15, bold=True, color=GREY, align=PP_ALIGN.CENTER)
rows = [  # colour, label, qr, link, bullets, tiles
    (YELLOW, "Talks", "qr-talks.png", "tele-task.de/series/1463", ["Guest talks on research and innovation"],
     ["tile-talk-1.png", "tile-talk-2.png", "tile-talk-3.jpeg"]),
    (ORANGE, "Work-\nshops", "qr-workshops.png", "aimaker.community",
     ["Hands-on topics", "For example: speech2summary, Docker for ML, semantic search"],
     ["tile-workshop-1.jpeg", "tile-workshop-2.jpeg", "tile-workshop-3.jpeg"]),
    (RED, "MOOCs", "qr-moocs.png", "open.hpi.de/channels/ai-service-center",
     ["ChatGPT: what does generative AI mean for society?", "Profitable AI", "Understanding and avoiding AI bias"],
     ["tile-mooc-1.jpeg", "tile-mooc-2.png", "tile-mooc-3.jpeg"]),
]
for i, (colour, label, code, link, bullets, tiles) in enumerate(rows):
    top = 0.55 + i * 3.6
    rect(s, 5.32, top, 1.75, 1.7, colour)
    text(s, 5.32, top, 1.75, 1.7, label, size=21, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=1.0)
    image(s, code, 9.0, top, w=1.5)
    text(s, 7.3, top + 1.55, 3.4, 0.35, link, size=13, bold=True, color="DC640D", align=PP_ALIGN.RIGHT)
    text(s, 5.32, top + 1.95, 5.6, 1.5, [[(f"•  {b}", {"space_after": 2})] for b in bullets], size=16)
    for j, tile in enumerate(tiles):
        image(s, tile, 11.0 + j * 3.0, top, w=2.82)
notes(s, "Everything is free and open. Newsletter via the QR code on the left.")

s = content("Infrastructure", DARK, None, logo=False)
image(s, "logo-kisz.png", 16.77, 0.43, w=2.8)
image(s, "qr-aisc.png", 1.6, 2.75, w=2.85)
text(s, 1.6, 5.65, 2.85, 0.35, "aisc.hpi.de · request access", size=13, color=GREY, align=PP_ALIGN.CENTER)
image(s, "compute-banner.jpeg", 5.09, 2.0, w=7.0)
image(s, "compute-racks.jpeg", 5.09, 5.95, w=7.0, h=5.2)
O = {"color": ORANGE}
terms = [[("•  Access ", {}), ("free of charge", O)], [("•  No production use", {})],
         [("      ◦  Data should be ", {}), ("anonymised or synthetic", O)], [("      ◦  No ", {}), ("hosting of products", O)],
         [("•  ", {}), ("Reporting & publication", O), (" by users", {})], [("•  ", {}), ("Existing rights", O), (" stay with users", {})],
         [("•  ", {}), ("New rights", O), (" stay with users", {})], [("      ◦  Usage rights granted for research and teaching", {})]]
text(s, 12.6, 1.95, 7.2, 4.2, terms, size=17, line=1.25)
specs = [("Training", ["64 NVIDIA H100 GPUs"]), ("Inference", ["40 NVIDIA A30 GPUs"]),
         ("ARM server", ["Ampere Altra Max M128-30 CPU", "2× NVIDIA L40 GPUs"]), ("GPU server", ["AMD Epyc CPU", "8× NVIDIA L40S GPUs"]),
         ("Edge", ["ARMv8 CPU", "NVIDIA Jetson AGX modules"]), ("Neuromorphic", ["288 SpiNNaker2 chips"]),
         ("Storage", ["1.5 PB NVRAM"]), ("Network", ["400 Gb/s InfiniBand", "200 Gb/s Ethernet"])]
for col, chunk in enumerate([specs[:4], specs[4:]]):
    paras = []
    for head, items in chunk:
        paras.append([(head, {"bold": True, "color": ORANGE})])
        paras += [[(f"•  {item}", {})] for item in items]
        paras.append([(" ", {"size": 6})])
    text(s, 12.6 + col * 3.7, 6.0, 3.6, 4.9, paras, size=16, line=1.1)
notes(s, "Free to use, no production use. Request access at aisc.hpi.de.")

s = content("Consulting", ORANGE, None, logo=False)
image(s, "qr-office-hours.png", 1.16, 2.55, w=2.95)
text(s, 1.16, 5.55, 2.95, 0.35, "Book office hours", size=13, bold=True, color=GREY, align=PP_ALIGN.CENTER)
blocks = [(RED, "AI office hours", ["AI infrastructure", "AI models & frameworks", "AI use cases"]),
          (ORANGE, "AI pilot projects", ["Co-developing a prototype", "Applications every three months",
                                         "Selection criteria such as AI maturity and public benefit", "Results are published"]),
          (YELLOW, "Collaborations", ["Jointly organised network meetings"]),
          (INK, "Past AI pilot projects", ["Generating maths problems", "Generating plain language", "Generating upcycling ideas",
                                           "Reducing food waste", "Dating documents by handwriting"])]
paras = []
for colour, head, items in blocks:
    paras.append([(head, {"bold": True, "color": colour, "size": 21})])
    paras += [[(f"•  {item}", {})] for item in items]
    paras.append([(" ", {"size": 8})])
text(s, 4.9, 2.2, 7.8, 8.8, paras, size=18, line=1.1)
image(s, "consulting.jpeg", 12.95, 0.3, w=6.6)
image(s, "qr-apply.png", 16.15, 7.0, w=2.95)
text(s, 16.15, 10.0, 2.95, 0.35, "Apply now!", size=13, bold=True, color="DC640D", align=PP_ALIGN.CENTER)
notes(s, "Concrete call to action: book office hours (QR on the left), apply for a pilot project, "
         "which opens every three months (QR on the right).")

s = content("RAG in practice", ORANGE, "IN PRACTICE")
for i, (shot, name, what, code) in enumerate([
        ("screenshot-sentra.png", "SENTRA", "Research Services of the German Bundestag: semantic search across "
         "expert reports, an AI summary, every statement traced to its report.", "github.com/aihpi/pilotproject-sentra"),
        ("screenshot-grundschutzki.png", "GrundschutzKI", "IT baseline protection: questions in plain language, "
         "a short answer with a jump link, the BSI original open next to it.", "github.com/aihpi/pilotprojekt-GrundschutzKI")]):
    x = 1.16 + i * 9.35
    pic = image(s, shot, x, 2.2, w=8.7)
    if pic.height > Inches(4.6):  # keep both screenshots the same height
        pic.height, pic.width = Inches(4.6), int(pic.width * Inches(4.6) / pic.height)
    text(s, x, 7.05, 8.7, 0.5, name, size=24, bold=True)
    text(s, x, 7.6, 8.5, 1.3, what, size=18, line=1.25)
    text(s, x, 8.95, 8.7, 0.4, code, size=15, color="DC640D", font="Courier New")
text(s, 1.16, 9.8, 17.5, 0.5, [[("Semantic search  ·  an AI summary  ·  sources for transparency.  ", {"bold": True}),
                                 ("Today you build the same thing, small.", {"color": GREY})]], size=19)
notes(s, "No live demo today, just these two to show RAG in real use. SENTRA: up to 2000 requests a year at the "
         "Bundestag's research services. GrundschutzKI: do-it-yourself instead of external consulting, also with NIS2 in mind.")

s = content("From document to chunks", BLUE, "BASICS")
kicker(s, 1.16, 2.0, 8, "Step 1 · Docling reads the page")
diagram("docling", 1.16, 2.5, 8.6, s)
text(s, 1.16, 6.3, 8.4, 1.4, "A PDF only stores where to draw each letter. Docling finds headings, paragraphs, tables and "
     "figures, and writes Markdown.", size=18, line=1.25)
kicker(s, 10.5, 2.0, 8, "Step 2 · cut into chunks")
diagram("chunking", 10.5, 2.5, 8.6, s)
text(s, 10.5, 6.3, 8.6, 4, [[("Our notebook: ", {"bold": True}), ("every 1000 characters, with 200 overlap. Simple, the same "
                             "for every document, but it may cut through a table.", {})],
                            [(" ", {"size": 8})],
                            [("GrundschutzKI: ", {"bold": True}), ("at every heading, because the BSI standards are clearly "
                             "structured. A chunk keeps its title. Docling ships several such chunkers.", {})]], size=18, line=1.25)
notes(s, "3 minutes. One method of many: fixed size with overlap is what the notebook does. "
         "Structured documents chunk better along headings.")

s = content("Prompt and citations", BLUE, "BASICS")
diagram("rag-prompt", 1.16, 2.0, 11.4, s)
diagram("citations", 1.16, 6.5, 9.3, s)
for i, (word, rest) in enumerate([("Retrieve", "the 5 closest chunks"), ("Augment", "the prompt with them, numbered [1] to [5]"),
                                  ("Generate", "the answer, citing those numbers")]):
    text(s, 13.4, 2.3 + i * 1.25, 6.2, 1.1, [[(word.upper(), {"bold": True, "color": BLUE, "spacing": 3, "size": 17})],
                                             [(rest, {"size": 19})]])
text(s, 13.4, 6.9, 6.0, 2.5, "Every [n] leads back to a chunk, and every chunk remembers its file and page. "
     "That is all a citation is.", size=19, line=1.25)
notes(s, "3 minutes. RAG = retrieve, augment, generate. Citations are bookkeeping, not magic.")

s = content("Is the answer any good?", BLUE, "EVALUATION")
text(s, 1.16, 2.0, 17.6, 0.9, "You scored faithfulness in notebook 4. Changed the chunk size or the number of chunks? Measure "
     "whether it helped. Neither score needs a hand-written correct answer.", size=20, line=1.25)
for i, (name, question, formula, how) in enumerate([
        ("Faithfulness", "Is every claim backed by the retrieved chunks?", "supported claims ÷ all claims",
         "A judge model splits the answer into claims and checks each one against the chunks. "
         "0.5 means half the claims have no source."),
        ("Relevance", "Does the answer address the question?", "similarity(questions from the answer, real question)",
         "A judge model writes the questions this answer would fit; their embeddings are compared with "
         "the question that was asked.")]):
    x = 1.16 + i * 9.0
    rect(s, x, 3.35, 8.6, 5.0, "F2F1EE")
    text(s, x + 0.4, 3.6, 7.9, 0.6, name, size=28, bold=True, color=BLUE)
    text(s, x + 0.4, 4.3, 7.9, 0.9, question, size=20, bold=True)
    rect(s, x + 0.4, 5.25, 7.8, 0.8, WHITE)
    text(s, x + 0.6, 5.25, 7.5, 0.8, formula, size=17, font="Courier New", anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.4, 6.35, 7.8, 1.9, how, size=17, line=1.25, color=GREY)
text(s, 1.16, 8.75, 17.6, 1.2, [[("In the template: ", {"bold": True}), ("both scores for every answer, a badge in the chat, "
     "and an evaluation app that compares configurations. Built on Ragas.", {})]], size=18, line=1.25)
notes(s, "1 to 2 minutes, a recap: they built faithfulness themselves in notebook 4. Relevance is new. Link it to the exercises: "
         "limit 2 versus 15, chunk size 300 versus 3000 are exactly the changes you would measure this way.")


# ---------------------------------------------------------------------------- page numbers, once the order is final
for number, slide in enumerate(prs.slides, start=1):
    if number > 1:
        text(slide, 19.3, 10.9, 0.5, 0.35, str(number), size=16, color=MUTED, align=PP_ALIGN.RIGHT)

prs.save(HERE / "rag-workshop.pptx")
print("wrote rag-workshop.pptx,", len(prs.slides), "slides")
